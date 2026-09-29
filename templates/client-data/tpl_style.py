"""Shared look for the consolidated client-data template.

The palette follows the Glow review the client likes (navy 002060 headers,
no gridlines, compact type) with the xlsx conventions for a workbook someone
fills in: yellow cells with blue text are typed, grey cells are calculated.
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
NAVY, NAVY_2, GOLD, GOLD_LT = "002060", "1F3864", "C3A962", "F3ECD8"
BAND, INPUT_BG, CALC_BG, LINE = "DCE3F1", "FFF2CC", "F2F2F2", "D9D9D9"
BLUE, RED, GREY, GREEN = "0000FF", "C00000", "595959", "006100"

F_MONEY = '#,##0;[Red](#,##0);"-"'
F_MONEY2 = '#,##0.00;[Red](#,##0.00);"-"'
F_QTY = '#,##0.000;[Red](#,##0.000);"-"'
F_RATE = '#,##0.00;[Red](#,##0.00);"-"'
F_NAV = '#,##0.0000;[Red](#,##0.0000);"-"'
F_PCT = '0.0%;[Red]-0.0%;"-"'
F_DATE = "dd-mmm-yyyy"
F_DAYS = '#,##0;[Red]-#,##0;"-"'
F_MULT = '0.00"x";[Red]-0.00"x";"-"'
F_TEXT = "@"

THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_LINE = Border(top=Side(style="thin", color=NAVY))
TOTAL_LINE = Border(top=Side(style="thin", color=NAVY), bottom=Side(style="double", color=NAVY))


def font(size=9, bold=False, italic=False, color="000000", underline=None):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color, underline=underline)


def fill(rgb):
    return PatternFill("solid", fgColor=rgb)


def col(n):
    return get_column_letter(n)


def letters(spec):
    """Map each column key in a register spec to its letter."""
    return {key: col(i + 1) for i, (key, *_rest) in enumerate(spec)}


def put(ws, ref, value, *, f=None, bg=None, fmt=None, al=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = f or font()
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    if al:
        c.alignment = al
    if border:
        c.border = border
    return c


LEFT = Alignment(horizontal="left", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def title_block(ws, title, subtitle, link_ref):
    """Row 1 title, row 2 subtitle, row 3 the DEMO banner; `link_ref` is where
    the way back to Start Here goes (clear of the title, which overflows right)."""
    ws.sheet_view.showGridLines = False
    put(ws, "A1", title, f=font(14, True, color=NAVY), al=LEFT)
    put(ws, "A2", subtitle, f=font(9, italic=True, color=GREY), al=LEFT)
    put(ws, "A3", "DEMO DATA — illustrative only: a fictional family, fictional securities and figures.",
        f=font(9, True, italic=True, color=RED), al=LEFT)
    put(ws, link_ref,
        '=HYPERLINK("#\'Start Here\'!A1","← Start Here")',
        f=font(9, color="0563C1", underline="single"), al=RIGHT)
    ws.row_dimensions[1].height = 22


def header(ws, row, start_col, labels, *, height=36, bg=NAVY):
    for i, text in enumerate(labels):
        put(ws, f"{col(start_col + i)}{row}", text, f=font(9, True, color="FFFFFF"),
            bg=bg, al=CENTER, border=BOX)
    ws.row_dimensions[row].height = height


def band(ws, row, c1, c2, text, bg=BAND, color=NAVY):
    ref = f"{col(c1)}{row}"
    if c2 > c1:
        ws.merge_cells(f"{ref}:{col(c2)}{row}")
    put(ws, ref, text, f=font(8, True, color=color), bg=bg, al=CENTER)
    for n in range(c1 + 1, c2 + 1):
        ws[f"{col(n)}{row}"].fill = fill(bg)


def widths(ws, mapping):
    for letter, w in mapping.items():
        ws.column_dimensions[letter].width = w


def list_validation(ws, cell_range, source, prompt=None):
    dv = DataValidation(type="list", formula1=source, allow_blank=True, showDropDown=False)
    dv.error = "Pick a value from the list (see the Lists sheet)."
    dv.errorTitle = "Not in the list"
    if prompt:
        dv.prompt, dv.promptTitle, dv.showInputMessage = prompt, "How to fill", True
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv


def define(wb, name, ref):
    dn = DefinedName(name, attr_text=ref)
    try:
        wb.defined_names.add(dn)
    except AttributeError:
        wb.defined_names[name] = dn
