"""Start Here: the settings every figure depends on, the overall check status,
a clickable index of the sheets, the colour legend and how to load a client.
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill

import tpl_demo as demo
from tpl_style import (BLUE, BOX, CALC_BG, CENTER, F_DATE, GOLD, GREY, INPUT_BG, LEFT,
                       LEFT_WRAP, NAVY, RED, band, define, fill, font, list_validation, put,
                       widths)

# (sheet, what it answers, kind)
INDEX = [
    ("Portfolio Allocation", "What the family owns by asset class and category: cost, value, "
     "allocation, gains and returns against the benchmark.", "View"),
    ("Basket Allocation", "Stable Growth, Entrepreneurial Growth, Liquidity, Thematic & "
     "Tactical against the family's targets.", "View"),
    ("Basket Detail", "Every holding in each basket, largest first, with its return.", "View"),
    ("Investor Summary", "Each entity's money by category, and its share of that entity's "
     "portfolio.", "View"),
    ("Period Change", "What moved since the previous valuation date: invested, redeemed, market "
     "gain, income.", "View"),
    ("Tax Summary", "Capital gains by taxpayer: book and taxable, long- and short-term, "
     "grandfathered; lots about to turn long-term.", "View"),
    ("Checks", "Does everything tie? Every view against the registers, every row check.",
     "Check"),
    ("Holdings", "One row per security per account: the position the views are built from.",
     "Register"),
    ("Tax Lots", "One row per purchase lot: purchase, sale, grandfathering, taxable gain.",
     "Register"),
    ("Transactions", "Dated money movements: buys, sells, SIPs, calls, distributions, income.",
     "Register"),
    ("Commitments", "Money promised to AIFs and private deals: committed, called, paid, "
     "uncalled, TVPI.", "Register"),
    ("Entities", "The family's taxpayers.", "Master data"),
    ("Accounts", "Where each holding sits, and the date of each provider's statement.",
     "Master data"),
    ("Securities", "What each holding is: category, basket, benchmark, price.", "Master data"),
    ("Lists", "The standard classification and every dropdown's values.", "Reference"),
    ("Column Mapping", "Where each column comes from in the two source formats, and how to "
     "convert it.", "Reference"),
]

STEPS = [
    "Set the client name, the valuation date and the previous valuation date above.",
    "Entities: one row per taxpayer. Accounts: one row per demat, PMS, folio, bank account or "
    "platform, with its statement date.",
    "Securities: one row per security or scheme. Pick its category from the dropdown — that "
    "sets asset class, basket and benchmark. Type the price at the valuation date.",
    "Holdings: one row per security per account. Type quantity and invested cost; the market "
    "value is quantity × price (or type the statement value in the override column).",
    "Tax Lots (listed shares and funds): one row per purchase lot, with any sale against it. "
    "Leave quantity, cost and realised gain on Holdings equal to the lots' totals.",
    "Transactions: every dated flow since the previous valuation date, at least. Commitments: "
    "every AIF or private deal with a drawdown.",
    "Open Checks. Every line must read OK before the views are shared. A red row check on a "
    "register names the problem on that row.",
    "Read the views. Change 'Show figures in' to switch between ₹ Crore, ₹ Lakh and ₹.",
    "Moving from another format? Column Mapping shows where each column comes from.",
]

CONVENTIONS = [
    "Registers hold full rupees. Views divide by the unit chosen above.",
    "Dates are real dates (dd-mmm-yyyy), never text. Percentages are fractions (0.125 = 12.5%).",
    "A blank means 'not reported'; 0 means a measured zero. Never type 0 for a figure you do "
    "not have.",
    "IDs join the sheets: E01 entity, A01 account, S01 security, H001 holding, L001 lot, T001 "
    "transaction, C01 commitment. Keep them unique and keep each list without gaps.",
    "Returns: holdings under a year show an absolute return; a year or more, the XIRR from the "
    "statement. Category and portfolio XIRRs are value-weighted and indicative.",
    "Private equity and pre-IPO shares are shown at cost until a priced round or a listing.",
    "Store PANs and account numbers masked. This file should never hold a full PAN.",
]


def build_start(wb):
    ws = wb.active
    ws.title = "Start Here"
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    put(ws, "B1", "Consolidated Client Data Template", f=font(16, True, color=NAVY), al=LEFT)
    put(ws, "B2", "One workbook for any client's portfolio: the registers hold the data in one "
        "standard shape, the views present it the way the family reads it.",
        f=font(9, italic=True, color=GREY), al=LEFT)
    put(ws, "B3", "DEMO DATA — illustrative only: a fictional family, fictional securities and "
        "figures.", f=font(9, True, italic=True, color=RED), al=LEFT)
    ws.row_dimensions[1].height = 26

    band(ws, 5, 2, 4, "SETTINGS", bg=NAVY, color="FFFFFF")
    settings = [
        (6, "Client / family name", demo.CLIENT, None, "ClientName",
         "Printed at the top of every view."),
        (7, "Valuation date", demo.VAL_DATE, F_DATE, "ValDate",
         "The date the market values are struck at."),
        (8, "Previous valuation date", demo.PREV_DATE, F_DATE, "PrevDate",
         "Start of the window the Period Change view measures."),
        (9, "Show figures in", demo.UNIT, None, "UnitLabel",
         "₹ Crore, ₹ Lakh or ₹. The registers always hold full rupees."),
        (11, "Grandfathering date (s.112A)", "31-Jan-2018", F_DATE, "GFDate",
         "Listed equity bought on or before this date is taxed from its FMV on this date."),
    ]
    from datetime import date
    for r, label, value, fmt, name, note in settings:
        put(ws, f"B{r}", label, f=font(9, True), al=LEFT, border=BOX)
        v = date(2018, 1, 31) if name == "GFDate" else value
        put(ws, f"C{r}", v, f=font(10, True, color=BLUE), bg=INPUT_BG, fmt=fmt, al=LEFT,
            border=BOX)
        put(ws, f"D{r}", note, f=font(8, italic=True, color=GREY), al=LEFT)
        define(wb, name, f"'Start Here'!$C${r}")
    put(ws, "B10", "Divide ₹ by", f=font(9, True), al=LEFT, border=BOX)
    put(ws, "C10", '=IFERROR(INDEX(L_UnitDiv,MATCH(UnitLabel,L_Unit,0)),1)', f=font(10, True),
        bg=CALC_BG, fmt="#,##0", al=LEFT, border=BOX)
    put(ws, "D10", "Set by 'Show figures in'; every view divides by it.",
        f=font(8, italic=True, color=GREY), al=LEFT)
    define(wb, "Divisor", "'Start Here'!$C$10")
    list_validation(ws, "C9", "L_Unit", "Pick the unit the views show.")

    put(ws, "B12", "Overall status", f=font(10, True, color=NAVY), al=LEFT, border=BOX)
    put(ws, "C12", "=OverallStatus", f=font(10, True), al=LEFT, border=BOX)
    put(ws, "D12", '=HYPERLINK("#\'Checks\'!A1","Open the Checks sheet →")',
        f=font(9, color="0563C1", underline="single"), al=LEFT)
    ws.conditional_formatting.add("C12", FormulaRule(
        formula=['LEFT($C$12,3)="ALL"'], font=Font(name="Arial", size=10, bold=True,
                                                    color="006100"),
        fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add("C12", FormulaRule(
        formula=['LEFT($C$12,3)<>"ALL"'], font=Font(name="Arial", size=10, bold=True,
                                                     color="9C0006"),
        fill=PatternFill("solid", fgColor="FFC7CE")))

    # Sheet index
    band(ws, 14, 2, 4, "WHAT IS IN THIS WORKBOOK — click a name to open it", bg=NAVY,
         color="FFFFFF")
    kind_bg = {"View": "F3ECD8", "Check": "E2EFDA", "Register": "FFF2CC",
               "Master data": "DDEBF7", "Reference": "EDEDED"}
    for i, (sheet, what, kind) in enumerate(INDEX):
        r = 15 + i
        put(ws, f"B{r}", f'=HYPERLINK("#\'{sheet}\'!A1","{sheet}")',
            f=font(9, True, color="0563C1", underline="single"), al=LEFT, border=BOX)
        put(ws, f"C{r}", kind, f=font(8, True, color=NAVY), bg=kind_bg[kind], al=CENTER,
            border=BOX)
        put(ws, f"D{r}", what, f=font(9), al=LEFT_WRAP, border=BOX)
        ws.row_dimensions[r].height = 24

    r0 = 15 + len(INDEX) + 1
    band(ws, r0, 2, 4, "HOW TO READ A SHEET", bg=NAVY, color="FFFFFF")
    legend = [
        ("Typed value", "Yellow cell, blue text: a figure or choice someone types.",
         INPUT_BG, BLUE),
        ("Calculated", "Grey cell, black text: a formula. Do not type over it.", CALC_BG,
         "000000"),
        ("Header", "Navy: column names. Gold rows on the views are group totals.", NAVY,
         "FFFFFF"),
        ("Row check", "Red when a row has a problem; green OK when it ties.", "FFC7CE",
         "9C0006"),
    ]
    for i, (label, text, bg, fg) in enumerate(legend):
        r = r0 + 1 + i
        put(ws, f"B{r}", label, f=font(9, True, color=fg), bg=bg, al=CENTER, border=BOX)
        put(ws, f"C{r}", text, f=font(9), al=LEFT)
        ws.merge_cells(f"C{r}:D{r}")

    r1 = r0 + len(legend) + 2
    band(ws, r1, 2, 4, "HOW TO LOAD A NEW CLIENT", bg=NAVY, color="FFFFFF")
    for i, step in enumerate(STEPS):
        r = r1 + 1 + i
        put(ws, f"B{r}", f"Step {i + 1}", f=font(9, True, color=NAVY), al=LEFT)
        put(ws, f"C{r}", step, f=font(9), al=LEFT_WRAP)
        ws.merge_cells(f"C{r}:D{r}")
        ws.row_dimensions[r].height = 26

    r2 = r1 + len(STEPS) + 2
    band(ws, r2, 2, 4, "CONVENTIONS", bg=NAVY, color="FFFFFF")
    for i, text in enumerate(CONVENTIONS):
        r = r2 + 1 + i
        put(ws, f"B{r}", "•", f=font(9, True, color=GOLD), al=Alignment_right())
        put(ws, f"C{r}", text, f=font(9), al=LEFT_WRAP)
        ws.merge_cells(f"C{r}:D{r}")
        ws.row_dimensions[r].height = 26

    widths(ws, {"A": 2, "B": 30, "C": 42, "D": 80})
    return ws


def Alignment_right():
    from openpyxl.styles import Alignment
    return Alignment(horizontal="right", vertical="top")
