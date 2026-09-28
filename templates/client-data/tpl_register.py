"""One writer for every register, so they all look and behave alike.

Layout: rows 1-3 title, subtitle and the DEMO banner (row 3 also carries the
totals of the visible rows, so they follow a filter); row 4 names the column
groups; row 5 is the header and the filter row; data starts on row 6 and every
formula is pre-extended to the register's last row.
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string

from tpl_schema import FIRST, HDR, LAST
from tpl_style import (BLUE, CALC_BG, CENTER, GREY, INPUT_BG, LEFT, NAVY, RIGHT, band, col,
                       fill, font, header, letters, list_validation, put, title_block)


def write_register(wb, sheet, spec, *, title, subtitle, link, rows, formulas, bands=(),
                   totals=(), validations=None, check_key="check", freeze="D6", tab=None,
                   wrap_header=True):
    ws = wb.create_sheet(sheet)
    L = letters(spec)
    last = LAST[sheet]
    title_block(ws, title, subtitle, link)
    if tab:
        ws.sheet_properties.tabColor = tab

    # Row 4: column groups
    for first_key, last_key, text in bands:
        c1 = column_index_from_string(L[first_key])
        c2 = column_index_from_string(L[last_key])
        band(ws, 4, c1, c2, text)

    # Row 5: headers (helpers get a quieter header)
    header(ws, HDR, 1, [h for _k, h, *_ in spec], height=48 if wrap_header else 30)
    for key, _h, kind, width, _fmt in spec:
        c = ws[f"{L[key]}{HDR}"]
        if kind == "h":
            c.fill = fill("808080")
        ws.column_dimensions[L[key]].width = width
        if kind == "h":
            ws.column_dimensions[L[key]].hidden = True

    # Row 3: totals of the rows a filter leaves visible
    if totals:
        first_total = min(column_index_from_string(L[k]) for k in totals)
        put(ws, f"{col(first_total - 1)}3", "Total of visible rows →",
            f=font(8, True, color=NAVY), al=RIGHT)
        for key in totals:
            fmt = next(f for k, _h, _kind, _w, f in spec if k == key)
            put(ws, f"{L[key]}3", f"=SUBTOTAL(9,{L[key]}{FIRST}:{L[key]}{last})",
                f=font(9, True, color=NAVY), fmt=fmt)

    # Data rows
    in_font, calc_font = font(9, color=BLUE), font(9)
    in_fill, calc_fill, help_fill = fill(INPUT_BG), fill(CALC_BG), fill("EDEDED")
    for r in range(FIRST, last + 1):
        data = rows[r - FIRST] if r - FIRST < len(rows) else {}
        for key, _h, kind, _w, fmt in spec:
            c = ws[f"{L[key]}{r}"]
            if kind == "in":
                v = data.get(key)
                if v not in (None, ""):
                    c.value = v
                c.font, c.fill = in_font, in_fill
            else:
                c.value = formulas[key].replace("{r}", str(r))
                c.font = calc_font
                c.fill = help_fill if kind == "h" else calc_fill
            if fmt:
                c.number_format = fmt

    # Dropdowns
    for key, (source, prompt) in (validations or {}).items():
        list_validation(ws, f"{L[key]}{FIRST}:{L[key]}{last}", source, prompt)

    # Row check: red when it is not OK
    if check_key in L:
        ref = f"{L[check_key]}{FIRST}:{L[check_key]}{last}"
        k = f"${L[check_key]}{FIRST}"
        ws.conditional_formatting.add(ref, FormulaRule(
            formula=[f'AND({k}<>"",{k}<>"OK")'],
            font=Font(name="Arial", size=9, bold=True, color="9C0006"),
            fill=PatternFill("solid", fgColor="FFC7CE")))
        ws.conditional_formatting.add(ref, FormulaRule(
            formula=[f'{k}="OK"'], font=Font(name="Arial", size=9, color="006100")))

    ws.freeze_panes = freeze
    last_col = L[spec[-1][0]]
    ws.auto_filter.ref = f"A{HDR}:{last_col}{last}"
    return ws, L

