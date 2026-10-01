"""Period Change (what moved between the previous valuation date and this one)
and Tax Summary (capital gains by taxpayer, from the Tax Lots register).

Period Change is Glow's month-on-month sheet made standard: every holding is
its own pot, so for each category
    value at the start + invested - redeemed + gain in value = value at the end
and dividend, interest and fees paid from outside sit beside the pot, never in
its value.

Tax Summary is the other client's lot-level tax register rolled up by
taxpayer: book profit against taxable gain (s.112A grandfathering), short and
long term, and the lots about to turn long-term.
"""
from openpyxl.worksheet.datavalidation import DataValidation

from tpl_schema import fx
from tpl_style import (BLUE, BOX, CENTER, F_DATE, F_DAYS, F_MONEY2, F_PCT, F_QTY, F_TEXT,
                       INPUT_BG, LEFT, NAVY, RED, RIGHT, band, define, font, put, widths)
from tpl_view_common import (category_grid, crit, grey_when_empty, note, page_setup, row_look,
                             total_look, view_title)
from tpl_view_pa import SPACER, _cell, _hdr

PC_FIRST, PC_TOTAL = 7, 31
TS_FIRST, TS_SLOTS = 11, 8
TS_TOTAL = TS_FIRST + TS_SLOTS           # 19
SOON_FIRST, SOON_SLOTS = 24, 10

D = "Divisor"
# a transaction counts in the period when dated after the previous valuation
# date and on or before this one
WIN = '[T.date],">"&PrevDate,[T.date],"<="&ValDate'


def build_period_change(wb):
    ws = wb.create_sheet("Period Change")
    view_title(ws, "Period Change",
               '="Figures in "&UnitLabel&" · from "&TEXT(PrevDate,"dd-mmm-yyyy")&" to "&'
               'TEXT(ValDate,"dd-mmm-yyyy")&" ("&(ValDate-PrevDate)&" days)"', "O1")

    _hdr(ws, "B5", "Category", "B6")
    _hdr(ws, "C5", '="Value on "&TEXT(PrevDate,"dd-mmm-yyyy")', "C6")
    _hdr(ws, "D5", "Money in and out", "E5")
    _hdr(ws, "D6", "Invested")
    _hdr(ws, "E6", "Redeemed / distributed")
    _hdr(ws, "F5", "Gain in value (market)", "F6")
    _hdr(ws, "G5", '="Value on "&TEXT(ValDate,"dd-mmm-yyyy")', "G6")
    _hdr(ws, "I5", "Outside the holding value", "J5")
    _hdr(ws, "I6", "Dividend / interest received (gross)")
    _hdr(ws, "J6", "Fees paid from outside")
    _hdr(ws, "L5", "Net gain", "L6")
    _hdr(ws, "M5", "Return for the period (%)", "M6")
    _hdr(ws, "O5", "Transactions in the period", "O6")
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 36

    def fill_row(r, k, kt):
        """One grid row; k / kt are the Holdings / Transactions criteria (or '')."""
        _cell(ws, f"C{r}", fx(f"=SUMIFS([H.prevmv]{k})/{D}" if k else f"=SUM([H.prevmv])/{D}"),
              F_MONEY2)
        _cell(ws, f"D{r}", fx(f'=-SUMIFS([T.net]{kt},[T.counts],"Investment",{WIN})/{D}'),
              F_MONEY2)
        _cell(ws, f"E{r}", fx(f'=SUMIFS([T.net]{kt},[T.counts],"Divestment",{WIN})/{D}'),
              F_MONEY2)
        _cell(ws, f"F{r}", f"=G{r}-C{r}-D{r}+E{r}", F_MONEY2)
        _cell(ws, f"G{r}", fx(f"=SUMIFS([H.mv]{k})/{D}" if k else f"=SUM([H.mv])/{D}"), F_MONEY2)
        _cell(ws, f"I{r}", fx(f'=SUMIFS([T.gross]{kt},[T.counts],"Income",{WIN})/{D}'), F_MONEY2)
        _cell(ws, f"J{r}", fx(f'=-SUMIFS([T.net]{kt},[T.counts],"Charges",{WIN})/{D}'), F_MONEY2)
        _cell(ws, f"L{r}", f"=F{r}+I{r}-J{r}", F_MONEY2)
        base = f"(C{r}+0.5*(D{r}-E{r}-I{r}))"
        _cell(ws, f"M{r}", f'=IF({base}<=0,"",L{r}/{base})', F_PCT)
        _cell(ws, f"O{r}", fx(f"=COUNTIFS({kt[1:]},{WIN})" if kt else
                              '=COUNTIFS([T.date],">"&PrevDate,[T.date],"<="&ValDate)'),
              "#,##0")

    for g in category_grid(PC_FIRST):
        r, kind = g["row"], g["kind"]
        put(ws, f"B{r}", g["label"], f=font(9), al=LEFT)
        fill_row(r, f",{crit(kind, 'H')},$B{r}", f",{crit(kind, 'T')},$B{r}")
        row_look(ws, r, kind, 2, 15)

    t = PC_TOTAL
    put(ws, f"B{t}", "Total", f=font(9, True, color=NAVY), al=LEFT)
    fill_row(t, "", "")
    total_look(ws, t, 2, 15)
    grey_when_empty(ws, f"B{PC_FIRST}:B{t - 2}", PC_FIRST, ["C", "D", "E", "G"])

    notes = [
        "How to read a row: value at the start + invested − redeemed or distributed + gain in "
        "value = value at the end. Every holding is its own pot; a flow moves money between the "
        "family and that pot.",
        "Dividend and interest are paid out of the holding, so they sit outside its value and "
        "are added to the gain (gross, before TDS). Fees charged inside a portfolio are already "
        "in its gain in value; fees the family pays from outside are shown separately and "
        "taken off.",
        "Money moved into or out of a bank account listed on Holdings is a Deposit or Withdrawal "
        "on that bank row, so it is counted once.",
        "Return for the period = net gain ÷ (value at the start + half of the money put in less "
        "the money taken out, income included) — a Simple Dietz return for this period only, "
        "not annualised.",
        "Values at the start come from each Holdings row's 'Market value at previous date'; "
        "transactions count when dated after the previous valuation date and on or before "
        "this one.",
    ]
    for i, text in enumerate(notes):
        note(ws, f"B{t + 2 + i}", text)

    widths(ws, {"A": 2, "B": 40, "C": 13, "D": 12, "E": 12, "F": 12, "G": 13, "H": SPACER,
                "I": 13, "J": 12, "K": SPACER, "L": 12, "M": 11, "N": SPACER, "O": 11})
    ws.freeze_panes = f"C{PC_FIRST}"
    page_setup(ws, f"A1:O{t + 2 + len(notes) - 1}")
    return ws


# (column, value key, term key, term, filtered by the chosen FY?)
TAX_COLS = [("D", "rbook", "rterm", "Short Term", True), ("E", "rbook", "rterm", "Long Term", True),
            ("F", "rtax", "rterm", "Short Term", True), ("G", "rtax", "rterm", "Long Term", True),
            ("H", "ubook", "uterm", "Short Term", False), ("I", "ubook", "uterm", "Long Term", False),
            ("J", "utax", "uterm", "Short Term", False), ("K", "utax", "uterm", "Long Term", False)]


def _tax_sum(value, term_key, term, by_fy, who):
    """SUMIFS over the lots; `who` is an extra criterion (the taxpayer) or ''."""
    base = f'[L.{value}]{who},[L.{term_key}],"{term}"'
    if not by_fy:
        return f"SUMIFS({base})"
    return f'IF(TaxFY="",SUMIFS({base}),SUMIFS({base},[L.fy],TaxFY))'


def build_tax_summary(wb):
    ws = wb.create_sheet("Tax Summary")
    view_title(ws, "Tax Summary",
               '="Figures in "&UnitLabel&" · realised: "&IF(TaxFY="","every year","FY "&TaxFY)'
               '&" · unrealised: at "&TEXT(ValDate,"dd-mmm-yyyy")', "N1")

    # Settings for this view
    put(ws, "B5", "Financial year of sale (realised gains)", f=font(9, True), al=LEFT, border=BOX)
    put(ws, "C5", "2025-26", f=font(10, True, color=BLUE), bg=INPUT_BG, fmt=F_TEXT, al=CENTER,
        border=BOX)
    note(ws, "D5", "Blank = every year. The unrealised columns are always at the valuation date.")
    put(ws, "B6", "Turning long-term within (days)", f=font(9, True), al=LEFT, border=BOX)
    put(ws, "C6", 90, f=font(10, True, color=BLUE), bg=INPUT_BG, fmt="#,##0", al=CENTER,
        border=BOX)
    note(ws, "D6", "Lists the lots whose long-term date falls within this many days of the "
         "valuation date.")
    define(wb, "TaxFY", "'Tax Summary'!$C$5")
    define(wb, "SoonDays", "'Tax Summary'!$C$6")

    dv = DataValidation(type="list", formula1="L_FY", allow_blank=True, showDropDown=False)
    dv.error, dv.errorTitle = "Pick a financial year from the list, or leave it blank.", "FY"
    dv.prompt, dv.promptTitle, dv.showInputMessage = (
        "The year of sale the realised columns show. Blank shows every year.", "Financial year",
        True)
    ws.add_data_validation(dv)
    dv.add("C5")
    dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="730",
                         allow_blank=False)
    dv2.error, dv2.errorTitle = "Type a whole number of days from 1 to 730.", "Days"
    ws.add_data_validation(dv2)
    dv2.add("C6")

    band(ws, 8, 2, 14, "CAPITAL GAINS BY TAXPAYER — from the Tax Lots sheet", bg=NAVY,
         color="FFFFFF")
    realised = '"Realised in "&IF(TaxFY="","any year","FY "&TaxFY)'
    _hdr(ws, "B9", "Taxpayer", "B10")
    _hdr(ws, "C9", "Lots on file", "C10")
    _hdr(ws, "D9", f'={realised}&" — book profit"', "E9")
    _hdr(ws, "F9", f'={realised}&" — taxable gain"', "G9")
    _hdr(ws, "H9", "Unrealised — book profit", "I9")
    _hdr(ws, "J9", "Unrealised — taxable gain", "K9")
    for letter in "DFHJ":
        _hdr(ws, f"{letter}10", "Short term")
    for letter in "EGIK":
        _hdr(ws, f"{letter}10", "Long term")
    _hdr(ws, "L9", "Dividend on lots (DPS × closing qty)", "L10")
    _hdr(ws, "M9", "Turning long-term soon", "N9")
    _hdr(ws, "M10", "Lots")
    _hdr(ws, "N10", "Unrealised taxable gain")
    ws.row_dimensions[9].height = 30
    ws.row_dimensions[10].height = 30

    t = TS_TOTAL
    for k in range(1, TS_SLOTS + 1):
        r = TS_FIRST + k - 1
        p = f"$P{r}"
        put(ws, f"P{r}", f'=IFERROR(INDEX(L_EntityID,{k})&"","")', f=font(8, color="808080"))
        put(ws, f"B{r}", fx(f'=IF({p}="","",IFERROR(INDEX([E.name],MATCH({p},[E.id],0))&"",{p}))'),
            f=font(9), al=LEFT)
        _cell(ws, f"C{r}", fx(f'=IF({p}="","",COUNTIFS([L.h_eid],{p}))'), "#,##0")
        who = f",[L.h_eid],{p}"
        for letter, value, term_key, term, by_fy in TAX_COLS:
            _cell(ws, f"{letter}{r}", fx(f'=IF(OR({p}="",N($C{r})=0),"",'
                                         f'{_tax_sum(value, term_key, term, by_fy, who)}/{D})'),
                  F_MONEY2)
        _cell(ws, f"L{r}", fx(f'=IF(OR({p}="",N($C{r})=0),"",SUMIFS([L.div],[L.h_eid],{p})/{D})'),
              F_MONEY2)
        _cell(ws, f"M{r}", fx(f'=IF({p}="","",COUNTIFS([L.h_eid],{p},[L.h_soon],1))'), "#,##0")
        _cell(ws, f"N{r}", fx(f'=IF(OR({p}="",N($M{r})=0),"",'
                              f'SUMIFS([L.utax],[L.h_eid],{p},[L.h_soon],1)/{D})'), F_MONEY2)

    put(ws, f"B{t}", "Total — every lot", f=font(9, True, color=NAVY), al=LEFT)
    _cell(ws, f"C{t}", fx("=COUNTA([L.lid])"), "#,##0")
    for letter, value, term_key, term, by_fy in TAX_COLS:
        _cell(ws, f"{letter}{t}", fx(f"={_tax_sum(value, term_key, term, by_fy, '')}/{D}"),
              F_MONEY2)
    _cell(ws, f"L{t}", fx(f"=SUM([L.div])/{D}"), F_MONEY2)
    _cell(ws, f"M{t}", fx("=COUNTIF([L.h_soon],1)"), "#,##0")
    _cell(ws, f"N{t}", fx(f"=SUMIFS([L.utax],[L.h_soon],1)/{D}"), F_MONEY2)
    total_look(ws, t, 2, 14)
    put(ws, f"B{t + 1}", fx(f'=IF(COUNTA([E.id])>{TS_SLOTS},COUNTA([E.id])-{TS_SLOTS}&'
                            f'" more taxpayer(s) are not listed here; the total row includes '
                            f'their lots."," ")'), f=font(8, True, italic=True, color=RED), al=LEFT)
    grey_when_empty(ws, f"B{TS_FIRST}:B{t - 1}", TS_FIRST, ["C"])

    # The lots about to turn long-term, soonest first
    band(ws, SOON_FIRST - 2, 2, 14, "", bg=NAVY, color="FFFFFF")
    ws[f"B{SOON_FIRST - 2}"].value = ('="LOTS TURNING LONG-TERM WITHIN "&SoonDays&" DAYS OF "&'
                                      'UPPER(TEXT(ValDate,"dd-mmm-yyyy"))&" — soonest first"')
    heads = [("B", "Security"), ("C", "Taxpayer"), ("D", "Lot"), ("E", "Purchase date"),
             ("F", "Closing quantity"), ("G", "Market value"), ("H", "Unrealised taxable gain"),
             ("I", "Becomes long-term on"), ("J", "Days left")]
    for letter, text in heads:
        _hdr(ws, f"{letter}{SOON_FIRST - 1}", text)
    ws.row_dimensions[SOON_FIRST - 1].height = 30
    for k in range(1, SOON_SLOTS + 1):
        r = SOON_FIRST + k - 1
        q = f"$P{r}"
        put(ws, f"P{r}", fx(f'=IFERROR(MATCH({k},[L.h_srank],0),"")'), f=font(8, color="808080"))

        def num(key, fmt, divide=False):
            v = f"INDEX([L.{key}],{q})"
            out = f"{v}/{D}" if divide else v
            _cell(ws, f"{letter}{r}", fx(f'=IF({q}="","",IF({v}="","",{out}))'), fmt)

        _cell(ws, f"B{r}", fx(f'=IF({q}="","",INDEX([L.name],{q})&"")'), None, LEFT)
        _cell(ws, f"C{r}", fx(f'=IF({q}="","",IFERROR(INDEX([E.short],MATCH(INDEX([L.h_eid],{q}),'
                              f'[E.id],0))&"",INDEX([L.ent],{q})&""))'), None, LEFT)
        _cell(ws, f"D{r}", fx(f'=IF({q}="","",INDEX([L.lid],{q})&"")'), None, CENTER)
        for letter, key, fmt, divide in [("E", "pdate", F_DATE, False), ("F", "cqty", F_QTY, False),
                                         ("G", "mval", F_MONEY2, True),
                                         ("H", "utax", F_MONEY2, True),
                                         ("I", "ltdate", F_DATE, False), ("J", "left", F_DAYS, False)]:
            num(key, fmt, divide)
        ws[f"E{r}"].alignment = CENTER
        ws[f"I{r}"].alignment = CENTER
    after = SOON_FIRST + SOON_SLOTS
    soon = fx("COUNTIF([L.h_soon],1)")
    put(ws, f"B{after}", f'=IF({soon}=0,"No lot turns long-term within "&SoonDays&" days of the '
                         f'valuation date.",IF({soon}>{SOON_SLOTS},{soon}-{SOON_SLOTS}&" more '
                         f'lot(s) turn long-term in this window: filter Tax Lots on \'Days left to '
                         f'become long-term\'."," "))',
        f=font(8, True, italic=True, color=RED), al=LEFT)

    notes = [
        "Book profit = sale (or market) value less purchase cost. Taxable gain uses the "
        "grandfathered cost for listed equity bought on or before 31-Jan-2018 (s.112A): the "
        "higher of the purchase price and the lower of the 31-Jan-2018 FMV and the sale (or "
        "market) price.",
        "Long-term after 12 months for listed shares, equity funds and listed units; after 24 "
        "months for unlisted shares and AIF units (Finance (No. 2) Act 2024). Debt funds bought "
        "from 1-Apr-2023 and market-linked debentures are never long-term (s.50AA). Each "
        "category's threshold is on the Lists sheet.",
        "No tax rate, surcharge, cess or the ₹1.25 lakh long-term exemption is applied: these "
        "are gains, not tax. Setting off losses is left to the tax adviser.",
        "Dividend on lots = the dividend per share typed on each lot × its closing quantity: an "
        "estimate for the current year, not what was received.",
        "Only the Tax Lots sheet feeds this view. A PMS or AIF reports its own capital gains on "
        "the manager's statement; they are not lots here.",
    ]
    n0 = after + 2
    for i, text in enumerate(notes):
        note(ws, f"B{n0 + i}", text)

    widths(ws, {"A": 2, "B": 32, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12,
                "I": 12, "J": 12, "K": 12, "L": 12, "M": 8, "N": 13, "O": 2})
    ws.column_dimensions["P"].hidden = True
    page_setup(ws, f"A1:N{n0 + len(notes) - 1}")
    return ws
