"""What every view shares: a title that names the client, the asset class ->
category grid, and the look of a group row, a category row and a total row.

Every label on the grid is a live reference to the Lists sheet, so renaming a
category on Lists renames it on every view.
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font

from tpl_style import (BAND, GOLD_LT, GREY, NAVY, RED, TOTAL_LINE, LEFT, RIGHT, fill, font,
                       col, put)
from tpl_taxonomy import ASSET_CLASSES, CATEGORIES, LIST_FIRST_ROW

VIEW_TAB = "C3A962"
INDENT = Alignment(horizontal="left", vertical="center", indent=1)
UNIT_LINE = '"Figures in "&UnitLabel&" · market values as at "&TEXT(ValDate,"dd-mmm-yyyy")'


def view_title(ws, title, subtitle_formula, link_ref):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = VIEW_TAB
    put(ws, "A1", f'="{title} — "&ClientName', f=font(14, True, color=NAVY), al=LEFT)
    put(ws, "A2", subtitle_formula, f=font(9, italic=True, color=GREY), al=LEFT)
    put(ws, "A3", "DEMO DATA — illustrative only: a fictional family, fictional securities "
        "and figures.", f=font(9, True, italic=True, color=RED), al=LEFT)
    put(ws, link_ref, '=HYPERLINK("#\'Start Here\'!A1","← Start Here")',
        f=font(9, color="0563C1", underline="single"), al=RIGHT)
    ws.row_dimensions[1].height = 22


def category_grid(first_row):
    """One gold row per asset class followed by its categories, in Lists order."""
    rows, r = [], first_row
    for i, ac in enumerate(ASSET_CLASSES):
        g = r
        rows.append({"row": r, "kind": "group", "group": g, "name": ac,
                     "label": f"=Lists!$A${LIST_FIRST_ROW + i}"})
        r += 1
        for j, (name, cat_ac, *_rest) in enumerate(CATEGORIES):
            if cat_ac == ac:
                rows.append({"row": r, "kind": "cat", "group": g, "name": name,
                             "label": f"=Lists!$D${LIST_FIRST_ROW + j}"})
                r += 1
    return rows


def crit(kind, reg):
    """The register column a grid row is matched on: asset class or category."""
    return f"[{reg}.ac]" if kind == "group" else f"[{reg}.cat]"


def group_rows(grid):
    return [g["row"] for g in grid if g["kind"] == "group"]


def cat_rows(grid):
    return [g["row"] for g in grid if g["kind"] == "cat"]


def row_look(ws, r, kind, c1, c2, label_col="B"):
    """Gold, bold group rows; indented category rows."""
    for n in range(c1, c2 + 1):
        c = ws[f"{col(n)}{r}"]
        if kind == "group":
            c.fill = fill(GOLD_LT)
            c.font = font(9, True, color=NAVY)
    if kind == "cat":
        ws[f"{label_col}{r}"].alignment = INDENT


def total_look(ws, r, c1, c2):
    for n in range(c1, c2 + 1):
        c = ws[f"{col(n)}{r}"]
        c.fill = fill(BAND)
        c.font = font(9, True, color=NAVY)
        c.border = TOTAL_LINE


def grey_when_empty(ws, label_range, first_row, value_cols):
    """Grey a row's label when every figure on it is zero: nothing is held there."""
    cond = "AND(" + ",".join(f"N(${c}{first_row})=0" for c in value_cols) + ")"
    ws.conditional_formatting.add(label_range, FormulaRule(
        formula=[cond], font=Font(name="Arial", size=9, color="A6A6A6")))


def page_setup(ws, area):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = area


def note(ws, ref, text):
    put(ws, ref, text, f=font(8, italic=True, color=GREY), al=LEFT)
