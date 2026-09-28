"""Checks: every view against the registers, and every register's own row check.

OK ties. ERROR must be fixed before the views are shared. INFO is worth reading
and is not wrong by itself. Money ties within one rupee, counts exactly. A
measure the sheet cannot compute reads "Could not calculate" and never passes.
"""
import math

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, PatternFill, Font

from tpl_schema import fx
from tpl_style import BAND, BOX, CENTER, GREY, NAVY, define, fill, font, header, put, widths
from tpl_view_basket import BA_FIRST, BA_TOTAL, BD_SECTIONS, BD_SLOTS
from tpl_view_common import cat_rows, category_grid, group_rows, note, page_setup, view_title
from tpl_view_pa import IS_FIRST, IS_TOTAL, PA_FIRST, PA_TOTAL
from tpl_view_period_tax import PC_FIRST, PC_TOTAL, TS_FIRST, TS_TOTAL

CHECK_TAB = "047857"
FIRST_ROW = 9
TOP_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
TOP_CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
PAIRS = ["C", "E", "G", "I", "K", "M", "O", "Q"]          # Investor Summary entity columns
SFX = 'IFERROR(INDEX(L_UnitSfx,MATCH(UnitLabel,L_Unit,0)),"")'


def q(sheet):
    return f"'{sheet}'" if (" " in sheet or "," in sheet) else sheet


def blocks(sheet, letter, rows):
    """Contiguous runs of `rows` as absolute ranges, comma-joined for SUM()."""
    rows, out = sorted(rows), []
    start = prev = rows[0]
    for r in rows[1:] + [None]:
        if r is not None and r == prev + 1:
            prev = r
            continue
        out.append(f"{q(sheet)}!${letter}${start}" + ("" if start == prev else f":${letter}${prev}"))
        if r is not None:
            start = prev = r
    return ",".join(out)


def gap(sheet, letter, rows, total):
    return f"ABS(SUM({blocks(sheet, letter, rows)})-N({q(sheet)}!${letter}${total}))"


class Book:
    def __init__(self, ws):
        self.ws, self.r, self.n = ws, FIRST_ROW, 0

    def group(self, title):
        ws, r = self.ws, self.r
        ws.merge_cells(f"B{r}:G{r}")
        put(ws, f"B{r}", title, f=font(9, True, color=NAVY), bg=BAND,
            al=Alignment(horizontal="left", vertical="center"))
        for c in "CDEFG":
            ws[f"{c}{r}"].fill = fill(BAND)
        ws.row_dimensions[r].height = 18
        self.r += 1

    def add(self, sheet, text, h, e, f, todo, e_len=40):
        ws, r = self.ws, self.r
        self.n += 1
        H = f"$H{r}"
        put(ws, f"B{r}", self.n, f=font(9, color=GREY), al=TOP_CENTER, border=BOX)
        put(ws, f"C{r}", f'=HYPERLINK("#{q(sheet)}!A1","{sheet}")',
            f=font(9, color="0563C1", underline="single"), al=TOP_WRAP, border=BOX)
        put(ws, f"D{r}", text, f=font(9), al=TOP_WRAP, border=BOX)
        put(ws, f"E{r}", fx("=" + e.replace("{h}", H)), f=font(9), al=TOP_WRAP, border=BOX)
        put(ws, f"F{r}", fx("=" + f.replace("{h}", H)), f=font(9, True), al=TOP_CENTER, border=BOX)
        put(ws, f"G{r}", todo, f=font(8, color=GREY), al=TOP_WRAP, border=BOX)
        put(ws, f"H{r}", fx("=" + h), f=font(8, color="808080"))
        lines = max(math.ceil(len(text) / 66), math.ceil(len(todo) / 78), math.ceil(e_len / 48), 1)
        ws.row_dimensions[r].height = max(15, 12 * lines + 4)
        self.r += 1

    # ---- kinds ---------------------------------------------------------------
    def rows(self, sheet, reg):
        rng = f"[{reg}.check]"
        n = f"SUMPRODUCT(--(LEN({rng})>0))"
        self.add(sheet, f"Every row on {sheet} passes its row check.",
                 f'IFERROR(SUMPRODUCT((LEN({rng})>0)*({rng}<>"OK")),-1)',
                 f'IF({{h}}<0,"Could not calculate: a row check shows an error",IF({n}=0,"No rows yet",'
                 f'IF({{h}}=0,{n}&" row(s), all OK",{{h}}&" of "&{n}&" row(s) need attention")))',
                 f'IF({{h}}<>0,"ERROR",IF({n}=0,"INFO","OK"))',
                 f"Filter the Row check column on {sheet} for anything but OK; each message names "
                 "the problem on that row.")

    def money(self, sheet, text, gaps, todo, gate=None, gate_text=""):
        m = f"IFERROR(MAX({','.join(gaps)})*Divisor,-1)"
        h = f"IF({gate},-2,{m})" if gate else m
        self.add(sheet, text, h,
                 f'IF({{h}}=-2,"{gate_text}",IF({{h}}<0,"Could not calculate",IF({{h}}<=1,"Ties",'
                 '"Off by ₹"&TEXT({h},"#,##0"))))',
                 'IF({h}=-2,"INFO",IF(AND({h}>=0,{h}<=1),"OK","ERROR"))', todo, 44)

    def count_tie(self, sheet, text, gaps, todo, gate=None, gate_text=""):
        m = f"IFERROR(MAX({','.join(gaps)}),-1)"
        h = f"IF({gate},-2,{m})" if gate else m
        self.add(sheet, text, h,
                 f'IF({{h}}=-2,"{gate_text}",IF({{h}}<0,"Could not calculate",IF({{h}}=0,"Ties",'
                 '"Off by "&{h})))',
                 'IF({h}=-2,"INFO",IF({h}=0,"OK","ERROR"))', todo, 44)

    def count(self, sheet, text, measure, ok_text, bad_expr, todo, level="INFO", e_len=80):
        self.add(sheet, text, f"IFERROR({measure},-1)",
                 f'IF({{h}}<0,"Could not calculate",IF({{h}}=0,"{ok_text}",{bad_expr}))',
                 f'IF({{h}}=0,"OK",IF({{h}}<0,"ERROR","{level}"))', todo, e_len)


def build_checks(wb):
    ws = wb.create_sheet("Checks")
    view_title(ws, "Checks", '="Every view against the registers, every row check · as at "&'
               'TEXT(ValDate,"dd-mmm-yyyy")', "G1")
    ws.sheet_properties.tabColor = CHECK_TAB
    header(ws, 8, 2, ["#", "Area", "Check", "Result", "Status", "What to do"], height=22)
    b = Book(ws)
    pa, ba, bd, inv, pc, ts = ("Portfolio Allocation", "Basket Allocation", "Basket Detail",
                               "Investor Summary", "Period Change", "Tax Summary")

    # ---- SETTINGS -----------------------------------------------------------------
    b.group("SETTINGS")
    b.add("Start Here", "The valuation date is later than the previous valuation date.",
          "IF(AND(ISNUMBER(ValDate),ISNUMBER(PrevDate)),IF(ValDate>PrevDate,0,1),1)",
          'IF(AND(ISNUMBER(ValDate),ISNUMBER(PrevDate)),TEXT(PrevDate,"dd-mmm-yyyy")&" to "&'
          'TEXT(ValDate,"dd-mmm-yyyy")&": "&(ValDate-PrevDate)&" days","Both must be real dates, '
          'the valuation date the later one")',
          'IF({h}=0,"OK","ERROR")',
          "Type both dates on Start Here. Period Change counts the flows between them.", 60)
    b.add("Start Here", "'Show figures in' is one of the units on the Lists sheet.",
          "IF(ISNUMBER(MATCH(UnitLabel,L_Unit,0)),0,1)",
          'IF({h}=0,"Views show "&UnitLabel&", every register figure divided by "&'
          'TEXT(Divisor,"#,##0"),"Not a unit on the Lists sheet: views divide by 1")',
          'IF({h}=0,"OK","ERROR")', "Pick the unit from the drop-down on Start Here.", 70)
    b.add(ts, "The tax year on Tax Summary is one of the years on the Lists sheet.",
          'IF(TaxFY="",-2,IF(ISNUMBER(MATCH(TaxFY,L_FY,0)),0,1))',
          'IF({h}=-2,"Blank: every year\'s sales are shown",IF({h}=0,"FY "&TaxFY,'
          '"Not a year on the Lists sheet"))',
          'IF({h}=-2,"INFO",IF({h}=0,"OK","ERROR"))',
          "Pick the year from the drop-down on Tax Summary, or leave it blank for every year.")
    b.add("Start Here", "The grandfathering date is 31-Jan-2018, the date the Finance Act 2018 fixes.",
          "IF(AND(ISNUMBER(GFDate),GFDate=DATE(2018,1,31)),0,1)",
          'IF({h}=0,"31-Jan-2018","Must be 31-Jan-2018")', 'IF({h}=0,"OK","ERROR")',
          "Put 31-Jan-2018 back on Start Here; Tax Lots grandfathers cost to the price on that date.")

    # ---- REGISTERS ----------------------------------------------------------------
    b.group("REGISTERS — every row's own check")
    for sheet, reg in [("Entities", "E"), ("Accounts", "A"), ("Securities", "S"), ("Holdings", "H"),
                       ("Tax Lots", "L"), ("Transactions", "T"), ("Commitments", "C")]:
        b.rows(sheet, reg)

    # ---- VIEWS TIE ----------------------------------------------------------------------
    b.group("VIEWS — each view adds up to its own total")
    tie = "Look for a holding with no category, asset class or basket; its row check on Holdings says which."
    pa_grid = category_grid(PA_FIRST)
    money_cols = ["D", "E", "I", "K", "M"]
    b.money(pa, "Category rows add to the Total row: cost, market value, income, gain incl. "
            "redeemed and gain on holdings.",
            [gap(pa, c, cat_rows(pa_grid), PA_TOTAL) for c in money_cols], tie)
    b.money(pa, "Asset-class rows add to the Total row, in the same five columns.",
            [gap(pa, c, group_rows(pa_grid), PA_TOTAL) for c in money_cols], tie)

    b.money(ba, "Basket rows add to the Total row: market value, cost and gain.",
            [gap(ba, c, range(BA_FIRST, BA_TOTAL), BA_TOTAL) for c in ["H", "N", "O"]], tie)
    b.money(ba, "Each basket's asset-class columns add to its own total.",
            [f"ABS(SUM({q(ba)}!$D${r}:$G${r})-N({q(ba)}!$H${r}))"
             for r in range(BA_FIRST, BA_TOTAL + 1)], tie)
    b.add(ba, "The target weights add to 100%.",
          f"IF(COUNT({q(ba)}!$K${BA_FIRST}:$K${BA_TOTAL - 1})=0,-2,"
          f"IF(ABS(N({q(ba)}!$K${BA_TOTAL})-1)>0.0001,1,0))",
          f'IF({{h}}=-2,"No targets typed",IF({{h}}=0,"100%","They add to "&'
          f'TEXT({q(ba)}!$K${BA_TOTAL},"0.0%")))',
          'IF({h}=-2,"INFO",IF({h}=0,"OK","ERROR"))',
          "Type the family's target for each basket on Basket Allocation; together they make 100%.")

    bd_tot = [s + 2 + BD_SLOTS for s in BD_SECTIONS]
    b.money(bd, "Each basket's total on Basket Detail equals its row on Basket Allocation (cost "
            "and market value).",
            [f"ABS(N({q(bd)}!${c}${t})-N({q(ba)}!${d}${BA_FIRST + i}))"
             for i, t in enumerate(bd_tot) for c, d in [("G", "N"), ("H", "H")]],
            "A basket code was renamed on the Lists sheet but not on Holdings or Securities.")
    listed = []
    for s, t in zip(BD_SECTIONS, bd_tot):
        cnt = fx(f"SUMPRODUCT(([H.basket]={q(bd)}!$R${s})*ISNUMBER([H.mv]))")
        listed.append(f"IF({cnt}>{BD_SLOTS},0,MAX("
                      f"ABS(SUM({q(bd)}!$G${s + 2}:$G${t - 1})-N({q(bd)}!$G${t})),"
                      f"ABS(SUM({q(bd)}!$H${s + 2}:$H${t - 1})-N({q(bd)}!$H${t}))))")
    b.money(bd, f"The holdings listed under each basket add to its total (a basket with more "
            f"than {BD_SLOTS} holdings is listed in part and counted under Capacity).", listed, tie)

    inv_grid = category_grid(IS_FIRST)
    ents = "COUNTA([E.id])>8"
    b.money(inv, "The entity columns add to the Family total.",
            [f"ABS(SUM({','.join(f'{q(inv)}!${v}${IS_TOTAL}' for v in PAIRS)})"
             f"-N({q(inv)}!$T${IS_TOTAL}))"], tie, gate=fx(ents),
            gate_text="More than 8 entities: the Family column includes ones not shown")
    b.money(inv, "Each entity's category rows, and its asset-class rows, add to its own total.",
            [gap(inv, v, rows, IS_TOTAL) for v in PAIRS
             for rows in (cat_rows(inv_grid), group_rows(inv_grid))], tie)
    b.money(inv, "The Family column's category rows, and its asset-class rows, add to its total.",
            [gap(inv, "T", cat_rows(inv_grid), IS_TOTAL), gap(inv, "T", group_rows(inv_grid), IS_TOTAL)],
            tie)

    pc_grid = category_grid(PC_FIRST)
    pc_cols = ["C", "D", "E", "G", "I", "J"]
    b.money(pc, "Category rows add to the Total row: start value, invested, redeemed, end value, "
            "income and fees.", [gap(pc, c, cat_rows(pc_grid), PC_TOTAL) for c in pc_cols], tie)
    b.money(pc, "Asset-class rows add to the Total row, in the same six columns.",
            [gap(pc, c, group_rows(pc_grid), PC_TOTAL) for c in pc_cols], tie)
    b.count_tie(pc, "The count of holdings adds up, by category and by asset class.",
                [gap(pc, "O", cat_rows(pc_grid), PC_TOTAL), gap(pc, "O", group_rows(pc_grid), PC_TOTAL)],
                tie)

    tx = "More than 8 taxpayers: the total includes ones not shown"
    trows = range(TS_FIRST, TS_TOTAL)
    b.money(ts, "Each taxpayer's gains, tax costs, dividends and near-term tax add to the Total row.",
            [gap(ts, c, trows, TS_TOTAL) for c in "DEFGHIJKLN"], tie, gate=fx(ents), gate_text=tx)
    b.count_tie(ts, "The counts of lots, and of lots turning long-term soon, add to the Total row.",
                [gap(ts, c, trows, TS_TOTAL) for c in "CM"], tie, gate=fx(ents), gate_text=tx)

    # ---- DATA COMPLETENESS -----------------------------------------------------------------
    b.group("DATA — what is missing or out of date")
    b.count("Holdings", "Every holding held at the previous valuation date has its value on that date.",
            "SUMPRODUCT(ISNUMBER([H.mv])*([H.prevmv]=\"\")*(1-ISNUMBER([H.first])*([H.first]>PrevDate)))",
            "Every holding has one",
            '{h}&" holding(s) have no previous value"',
            "Type the previous statement's value on Holdings. Without it Period Change counts the "
            "whole value as gain.", level="ERROR")
    b.count("Holdings", "Every holding with a return shown as XIRR has its XIRR typed.",
            'COUNTIF([H.shown],"XIRR missing")', "None missing",
            '{h}&" holding(s) show XIRR missing"',
            "Type the XIRR from the statement, or leave it: the view falls back to the absolute return.")
    b.count("Holdings", "Every holding has a first investment date.", 'COUNTIF([H.shown],"No date")',
            "All dated", '{h}&" holding(s) have no first date, so no return is shown"',
            "A bank balance needs none. For anything else type the first purchase date on Holdings.")
    b.count("Accounts", "Every account's statement is dated at the valuation date.",
            "SUMPRODUCT(ISNUMBER([A.asof])*([A.asof]<ValDate))", "All current",
            '{h}&" account(s) are on an older statement"',
            "Their values are as of their own date. Load the latest statement when it arrives.")
    b.count("Securities", "Every price is dated at the valuation date.",
            "SUMPRODUCT(ISNUMBER([S.pdate])*([S.pdate]<ValDate))", "All current",
            '{h}&" price(s) are older than the valuation date"',
            "Usually a fund that reports monthly or quarterly. Update the price when it is published.")
    b.count("Transactions", "Transactions dated on or before the previous valuation date.",
            'COUNTIFS([T.date],"<="&PrevDate)', "None",
            '{h}&" transaction(s): kept for the record, not counted in Period Change"',
            "Nothing to do; Period Change counts only the flows after the previous valuation date.")
    b.count("Commitments", "Money called and not yet paid.", 'COUNTIF([C.unpaid],">0")', "None unpaid",
            '{h}&" commitment(s) have money called and not yet paid: ₹"&TEXT(SUMIF([C.unpaid],">0")'
            '/Divisor,"#,##0.00")&IF(' + SFX + '="₹",""," "&' + SFX + ')',
            "Pay the call, or type the payment on Commitments once it is made.", e_len=110)
    b.count(ts, "Lots turning long-term soon.", "COUNTIF([L.h_soon],1)", "None",
            '{h}&" lot(s) turn long-term within "&SoonDays&" days: see Tax Summary"',
            "Selling one before its long-term date is taxed as short-term.")

    # ---- CAPACITY -------------------------------------------------------------------------------
    b.group("CAPACITY — rows the views have room for")
    grow = "Add rows to the view by copying its last row down, then widen this check."
    b.count(ba, "Baskets beyond the 4 the basket views show.", "MAX(0,COUNTA(Lists!$K$6:$K$65)-4)",
            "Within 4", '{h}&" basket(s) not shown on the basket views"', grow)
    b.count(inv, "Entities beyond the 8 columns on Investor Summary.", "MAX(0,COUNTA([E.id])-8)",
            "Within 8", '{h}&" entit(ies) not given a column"', grow)
    b.count(pa, "Commitments beyond the 10 the Portfolio Allocation block lists.",
            "MAX(0,COUNTA([C.cid])-10)", "Within 10", '{h}&" commitment(s) not listed"', grow)
    over = "+".join(fx(f"(SUMPRODUCT(([H.basket]={q(ba)}!$C${r})*ISNUMBER([H.mv]))>{BD_SLOTS})")
                    for r in range(BA_FIRST, BA_TOTAL))
    b.count(bd, f"Baskets holding more than the {BD_SLOTS} rows Basket Detail lists.", over,
            f"Within {BD_SLOTS}", '{h}&" basket(s) listed in part"', grow)
    b.count(pa, "Categories beyond the 19 the views show.", "MAX(0,COUNTA(Lists!$D$6:$D$65)-19)",
            "Within 19", '{h}&" categor(ies) not shown"', grow)
    b.count(pa, "Asset classes beyond the 4 the views show.", "MAX(0,COUNTA(Lists!$A$6:$A$65)-4)",
            "Within 4", '{h}&" asset class(es) not shown"', grow)
    regs = [("Entities", "E.id", 20, 18), ("Accounts", "A.id", 50, 45), ("Securities", "S.id", 100, 90),
            ("Holdings", "H.hid", 200, 180), ("Tax Lots", "L.lid", 200, 180),
            ("Transactions", "T.tid", 300, 270), ("Commitments", "C.cid", 30, 27)]
    measure = "+".join(f"(COUNTA([{k}])>={th})" for _s, k, _n, th in regs)
    names = "&".join(f'IF(COUNTA([{k}])>={th},"; {s} ("&COUNTA([{k}])&" of {n})","")'
                     for s, k, n, th in regs)
    b.count("Start Here", "Registers 90% or more full.", measure, "All have room",
            f"MID({names},3,400)",
            "Insert rows above a register's last row, so every formula's range grows with it, and "
            "copy the formulas down.")

    # ---- overall -----------------------------------------------------------------------------------
    last = b.r - 1
    st = f"$F${FIRST_ROW}:$F${last}"
    ws.merge_cells("B5:C5")
    ws.merge_cells("D5:G5")
    ws.merge_cells("B6:C6")
    ws.merge_cells("D6:G6")
    put(ws, "B5", "Overall status", f=font(10, True, color=NAVY), al=CENTER, border=BOX)
    put(ws, "D5", f'=IF(COUNTIF({st},"ERROR")=0,"ALL CHECKS PASS"&IF(COUNTIF({st},"INFO")>0," · "&'
                  f'COUNTIF({st},"INFO")&" note(s) to read",""),COUNTIF({st},"ERROR")&'
                  f'" CHECK(S) NEED ATTENTION")', f=font(11, True), al=CENTER, border=BOX)
    put(ws, "B6", "Lines", f=font(9, color=GREY), al=CENTER)
    put(ws, "D6", f'=COUNTIF({st},"OK")&" OK · "&COUNTIF({st},"INFO")&" note(s) · "&'
                  f'COUNTIF({st},"ERROR")&" error(s)"', f=font(9, color=GREY), al=CENTER)
    ws.row_dimensions[5].height = 24
    define(wb, "OverallStatus", "'Checks'!$D$5")

    green = (PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE"), Font(color="006100", bold=True))
    red = (PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE"), Font(color="9C0006", bold=True))
    amber = (PatternFill("solid", start_color="FFEB9C", end_color="FFEB9C"), Font(color="9C5700", bold=True))
    cf = ws.conditional_formatting
    cf.add("D5", FormulaRule(formula=['LEFT($D$5,3)="ALL"'], fill=green[0], font=green[1]))
    cf.add("D5", FormulaRule(formula=['LEFT($D$5,3)<>"ALL"'], fill=red[0], font=red[1]))
    for word, (fl, fo) in [("OK", green), ("ERROR", red), ("INFO", amber)]:
        cf.add(f"F{FIRST_ROW}:F{last}",
               FormulaRule(formula=[f'$F{FIRST_ROW}="{word}"'], fill=fl, font=fo))

    note(ws, f"B{last + 2}", "OK: ties · ERROR: fix before the views are shared · INFO: worth "
         "reading, not wrong by itself. Money ties within ₹1; counts exactly.")
    widths(ws, {"A": 2, "B": 5, "C": 22, "D": 60, "E": 44, "F": 10, "G": 60, "H": 8})
    ws.column_dimensions["H"].hidden = True
    ws.freeze_panes = "A9"
    page_setup(ws, f"A1:G{last + 3}")
    return ws
