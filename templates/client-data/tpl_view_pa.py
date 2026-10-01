"""Portfolio Allocation (Glow's own layout: category rows under gold asset-class
rows, cost against value, allocation, gains, and returns against a blended
index) and the Investor Summary matrix (each entity's money by category).

Every figure is a SUMIFS on the Holdings register, divided by the display unit,
so the views never hold a typed number.
"""
from tpl_schema import fx
from tpl_style import (BOX, CENTER, F_DATE, F_MONEY2, F_MULT, F_PCT, LEFT, NAVY, RIGHT, band,
                       font, put, widths)
from tpl_view_common import (UNIT_LINE, cat_rows, category_grid, crit, grey_when_empty,
                             group_rows, note, page_setup, row_look, total_look, view_title)

PA_FIRST, PA_TOTAL = 7, 31
IS_FIRST, IS_TOTAL = 8, 32
SPACER = 1.7


def _hdr(ws, ref, text, merge_to=None):
    put(ws, ref, text, f=font(9, True, color="FFFFFF"), bg=NAVY, al=CENTER, border=BOX)
    if merge_to:
        ws.merge_cells(f"{ref}:{merge_to}")
        # fill every cell of the merge so the border and colour run the whole width
        from openpyxl.utils import range_boundaries
        c1, r1, c2, r2 = range_boundaries(f"{ref}:{merge_to}")
        from tpl_style import col, fill
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                cell = ws[f"{col(cc)}{rr}"]
                cell.fill, cell.border = fill(NAVY), BOX


def _cell(ws, ref, formula, fmt, al=RIGHT):
    put(ws, ref, formula, f=font(9), fmt=fmt, al=al)


def build_portfolio_allocation(wb):
    ws = wb.create_sheet("Portfolio Allocation")
    view_title(ws, "Portfolio Allocation", "=" + UNIT_LINE, "T1")

    # Two-row header, Glow's layout: spacer columns H, J, L, N and Q
    for ref, text, to in [("B5", "Category", "B6"), ("C5", "Investment date range", "C6"),
                          ("D5", "Investment at cost", "D6"), ("E5", "Market value", "E6"),
                          ("F5", "Allocation (% of portfolio)", "F6"),
                          ("G5", "Allocation (% of asset class)", "G6"),
                          ("I5", "Dividend / interest", "I6"),
                          ("K5", "Gain / (loss) including redeemed", "K6"),
                          ("M5", "Gain / (loss) on what is held", "M6"),
                          ("O5", "Absolute return (%)", "P5"),
                          ("R5", "Annualised return — XIRR (%)", "T5")]:
        _hdr(ws, ref, text, to)
    for ref, text in [("O6", "Portfolio"), ("P6", "Blended index"), ("R6", "Portfolio"),
                      ("S6", "Blended index"), ("T6", "Share of value covered")]:
        _hdr(ws, ref, text)
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 30

    grid = category_grid(PA_FIRST)
    groups = {g["row"]: [c["row"] for c in grid if c["kind"] == "cat" and c["group"] == g["row"]]
              for g in grid if g["kind"] == "group"}
    D = "Divisor"
    for g in grid:
        r, kind = g["row"], g["kind"]
        k = crit(kind, "H")
        put(ws, f"B{r}", g["label"], f=font(9), al=LEFT)
        _cell(ws, f"C{r}", fx(
            f'=IF(COUNTIFS({k},$B{r},[H.first],">0")=0,"",'
            f'TEXT(_xlfn.MINIFS([H.first],{k},$B{r}),"mmm-yy")&" – "&'
            f'TEXT(MAX(_xlfn.MAXIFS([H.first],{k},$B{r}),_xlfn.MAXIFS([H.last],{k},$B{r})),'
            f'"mmm-yy"))'), None, CENTER)
        _cell(ws, f"D{r}", fx(f"=SUMIFS([H.cost],{k},$B{r})/{D}"), F_MONEY2)
        _cell(ws, f"E{r}", fx(f"=SUMIFS([H.mv],{k},$B{r})/{D}"), F_MONEY2)
        _cell(ws, f"F{r}", f'=IF(N($E${PA_TOTAL})=0,"",E{r}/$E${PA_TOTAL})', F_PCT)
        if kind == "cat":
            _cell(ws, f"G{r}", f'=IF(N($E${g["group"]})=0,"",E{r}/$E${g["group"]})', F_PCT)
        _cell(ws, f"I{r}", fx(f"=SUMIFS([H.income],{k},$B{r})/{D}"), F_MONEY2)
        _cell(ws, f"K{r}", fx(f"=(SUMIFS([H.rgain],{k},$B{r})+SUMIFS([H.ugain],{k},$B{r}))/{D}"),
              F_MONEY2)
        _cell(ws, f"M{r}", fx(f"=SUMIFS([H.ugain],{k},$B{r})/{D}"), F_MONEY2)
        _cell(ws, f"O{r}", f'=IF(N(D{r})=0,"",M{r}/D{r})', F_PCT)
        _cell(ws, f"P{r}", fx(f'=IF(SUMIFS([H.h_mbam],{k},$B{r})=0,"",'
                              f'SUMIFS([H.h_mba],{k},$B{r})/SUMIFS([H.h_mbam],{k},$B{r}))'), F_PCT)
        _cell(ws, f"R{r}", fx(f'=IF(SUMIFS([H.h_m],{k},$B{r})=0,"",'
                              f'SUMIFS([H.h_mx],{k},$B{r})/SUMIFS([H.h_m],{k},$B{r}))'), F_PCT)
        _cell(ws, f"S{r}", fx(f'=IF(SUMIFS([H.h_m],{k},$B{r})=0,"",'
                              f'SUMIFS([H.h_mbx],{k},$B{r})/SUMIFS([H.h_m],{k},$B{r}))'), F_PCT)
        _cell(ws, f"T{r}", fx(f'=IF(N(E{r})=0,"",SUMIFS([H.h_m],{k},$B{r})/{D}/E{r})'), F_PCT)
        row_look(ws, r, kind, 2, 20)

    # Grand total: straight from the register, so Checks can tie the rows to it
    t = PA_TOTAL
    put(ws, f"B{t}", "Total", f=font(9, True, color=NAVY), al=LEFT)
    _cell(ws, f"C{t}", fx('=IF(COUNT([H.first])=0,"",TEXT(MIN([H.first]),"mmm-yy")&" – "&'
                          'TEXT(MAX(MAX([H.first]),MAX([H.last])),"mmm-yy"))'), None, CENTER)
    _cell(ws, f"D{t}", fx(f"=SUM([H.cost])/{D}"), F_MONEY2)
    _cell(ws, f"E{t}", fx(f"=SUM([H.mv])/{D}"), F_MONEY2)
    _cell(ws, f"F{t}", "=" + "+".join(f"N(F{r})" for r in groups), F_PCT)
    _cell(ws, f"I{t}", fx(f"=SUM([H.income])/{D}"), F_MONEY2)
    _cell(ws, f"K{t}", fx(f"=(SUM([H.rgain])+SUM([H.ugain]))/{D}"), F_MONEY2)
    _cell(ws, f"M{t}", fx(f"=SUM([H.ugain])/{D}"), F_MONEY2)
    _cell(ws, f"O{t}", f'=IF(N(D{t})=0,"",M{t}/D{t})', F_PCT)
    _cell(ws, f"P{t}", fx('=IF(SUM([H.h_mbam])=0,"",SUM([H.h_mba])/SUM([H.h_mbam]))'), F_PCT)
    _cell(ws, f"R{t}", fx('=IF(SUM([H.h_m])=0,"",SUM([H.h_mx])/SUM([H.h_m]))'), F_PCT)
    _cell(ws, f"S{t}", fx('=IF(SUM([H.h_m])=0,"",SUM([H.h_mbx])/SUM([H.h_m]))'), F_PCT)
    _cell(ws, f"T{t}", fx(f'=IF(N(E{t})=0,"",SUM([H.h_m])/{D}/E{t})'), F_PCT)
    total_look(ws, t, 2, 20)
    grey_when_empty(ws, f"B{PA_FIRST}:B{PA_TOTAL - 2}", PA_FIRST, ["D", "E"])

    note(ws, f"B{t + 2}", "Returns: a holding held under a year shows its absolute return; a year "
         "or more, the XIRR its statement reports. A category's XIRR is value-weighted across its "
         "holdings, so read it as indicative.")
    note(ws, f"B{t + 3}", "Private equity and pre-IPO shares are shown at cost until a priced round "
         "or a listing, so they carry no gain. Cash balances carry no return.")
    note(ws, f"B{t + 4}", "Blended index: the benchmark return over the holdings that report one — "
         "cost-weighted for the absolute return, value-weighted for the XIRR. 'Share of value "
         "covered' is the part of the row's value that has an XIRR.")
    note(ws, f"B{t + 5}", "Gain including redeemed = gain on what is held + realised gain on exits "
         "(Holdings). Dividend / interest is what each statement reports as received.")

    _commitments_block(ws, t + 7)

    widths(ws, {"A": 2, "B": 40, "C": 16, "D": 12, "E": 12, "F": 10, "G": 10, "H": SPACER,
                "I": 11, "J": SPACER, "K": 13, "L": SPACER, "M": 13, "N": SPACER, "O": 10,
                "P": 10, "Q": SPACER, "R": 10, "S": 10, "T": 12})
    ws.freeze_panes = f"C{PA_FIRST}"
    page_setup(ws, f"A1:T{t + 21}")
    return ws


def _commitments_block(ws, top):
    """Glow's review has no view of drawdown commitments; this adds one, the first
    ten from the Commitments register (the total row covers every one)."""
    band(ws, top, 2, 20, "PRIVATE MARKET COMMITMENTS — AIFs and private deals with a drawdown "
                        "(first 10 rows of the Commitments sheet; the total covers every row)",
         bg=NAVY, color="FFFFFF")
    hdr = top + 1
    cols = [("B", "Fund", "fund", "text"), ("C", "Entity", "ent", "text"),
            ("D", "Committed", "comm", "money"), ("E", "Called to date", "called", "money"),
            ("F", "% called", "pct", "pct"), ("G", "Paid to date", "paid", "money"),
            ("I", "Uncalled", "uncalled", "money"), ("K", "Called, not yet paid", "unpaid", "money"),
            ("M", "Distributions received", "dist", "money"),
            ("O", "Current value", "value", "money"), ("P", "TVPI (x)", "tvpi", "mult"),
            ("R", "DPI (x)", "dpi", "mult"), ("S", "Next call date", "ncdate", "date"),
            ("T", "Next call amount", "ncamt", "money")]
    for letter, text, _k, _t in cols:
        _hdr(ws, f"{letter}{hdr}", text)
    ws.row_dimensions[hdr].height = 30
    fmts = {"money": F_MONEY2, "pct": F_PCT, "mult": F_MULT, "date": F_DATE, "text": None}
    first = hdr + 1
    for i in range(10):
        r, k = first + i, i + 1
        empty = f"INDEX([C.cid],{k})=\"\""
        for letter, _text, key, kind in cols:
            val = f"INDEX([C.{key}],{k})"
            if kind == "text":
                frm = f'=IF({empty},"",{val}&"")'
            elif kind == "money":
                frm = f'=IF({empty},"",IF({val}="","",{val}/Divisor))'
            else:
                frm = f'=IF({empty},"",IF({val}="","",{val}))'
            put(ws, f"{letter}{r}", fx(frm), f=font(9), fmt=fmts[kind],
                al=LEFT if kind == "text" else RIGHT)
    tot = first + 10
    put(ws, f"B{tot}", "Total — every commitment", f=font(9, True, color=NAVY), al=LEFT)
    for letter, key in [("D", "comm"), ("E", "called"), ("G", "paid"), ("I", "uncalled"),
                        ("K", "unpaid"), ("M", "dist"), ("O", "value"), ("T", "ncamt")]:
        put(ws, f"{letter}{tot}", fx(f"=SUM([C.{key}])/Divisor"), f=font(9), fmt=F_MONEY2,
            al=RIGHT)
    put(ws, f"F{tot}", f'=IF(N(D{tot})=0,"",E{tot}/D{tot})', f=font(9), fmt=F_PCT, al=RIGHT)
    put(ws, f"P{tot}", f'=IF(N(G{tot})=0,"",(O{tot}+M{tot})/G{tot})', f=font(9), fmt=F_MULT,
        al=RIGHT)
    put(ws, f"R{tot}", f'=IF(N(G{tot})=0,"",M{tot}/G{tot})', f=font(9), fmt=F_MULT, al=RIGHT)
    put(ws, f"S{tot}", fx('=IF(COUNT([C.ncdate])=0,"",MIN([C.ncdate]))'), f=font(9),
        fmt=F_DATE, al=RIGHT)
    total_look(ws, tot, 2, 20)
    put(ws, f"B{tot + 1}", fx('=IF(COUNTA([C.cid])>10,COUNTA([C.cid])-10&" more commitment(s): '
                              'see the Commitments sheet. The total row includes them.","")'),
        f=font(8, True, italic=True, color="C00000"), al=LEFT)
    note(ws, f"B{tot + 2}", "TVPI = (current value + distributions) ÷ paid; DPI = distributions ÷ "
         "paid. The total's next call date is the earliest one. Current value is the Holdings "
         "row for the same account and security.")


def build_investor_summary(wb):
    ws = wb.create_sheet("Investor Summary")
    view_title(ws, "Investor Summary", "=" + UNIT_LINE + '&" · each entity\'s column: its value '
               'by category and the share of its own portfolio"', "U1")
    pairs = ["C", "E", "G", "I", "K", "M", "O", "Q"]
    from tpl_style import col
    from openpyxl.utils import column_index_from_string as ci

    _hdr(ws, "B6", "Category", "B7")
    for i, v in enumerate(pairs):
        s = col(ci(v) + 1)
        put(ws, f"{v}5", f"=IFERROR(INDEX(L_EntityID,{i + 1}),\"\")", f=font(8, color="808080"))
        _hdr(ws, f"{v}6", fx(f'=IF({v}$5="","",IFERROR(INDEX([E.short],MATCH({v}$5,[E.id],0))&"",'
                             f'{v}$5))'), f"{s}6")
        _hdr(ws, f"{v}7", "Value")
        _hdr(ws, f"{s}7", "% of theirs")
    _hdr(ws, "T6", "Family", "U6")
    _hdr(ws, "T7", "Value")
    _hdr(ws, "U7", "% of family")
    ws.row_dimensions[5].hidden = True
    ws.row_dimensions[6].height = 24
    ws.row_dimensions[7].height = 24

    t = IS_TOTAL
    grid = category_grid(IS_FIRST)
    for g in grid:
        r, kind = g["row"], g["kind"]
        k = crit(kind, "H")
        put(ws, f"B{r}", g["label"], f=font(9), al=LEFT)
        for v in pairs:
            s = col(ci(v) + 1)
            _cell(ws, f"{v}{r}", fx(f'=IF({v}$5="","",SUMIFS([H.mv],[H.h_eid],{v}$5,{k},$B{r})'
                                    f'/Divisor)'), F_MONEY2)
            _cell(ws, f"{s}{r}", f'=IF(OR({v}$5="",N({v}${t})=0),"",{v}{r}/{v}${t})', F_PCT)
        _cell(ws, f"T{r}", fx(f"=SUMIFS([H.mv],{k},$B{r})/Divisor"), F_MONEY2)
        _cell(ws, f"U{r}", f'=IF(N($T${t})=0,"",T{r}/$T${t})', F_PCT)
        row_look(ws, r, kind, 2, 21)

    put(ws, f"B{t}", "Total", f=font(9, True, color=NAVY), al=LEFT)
    for v in pairs:
        s = col(ci(v) + 1)
        _cell(ws, f"{v}{t}", fx(f'=IF({v}$5="","",SUMIFS([H.mv],[H.h_eid],{v}$5)/Divisor)'),
              F_MONEY2)
        _cell(ws, f"{s}{t}", f'=IF(OR({v}$5="",N($T${t})=0),"",{v}{t}/$T${t})', F_PCT)
    _cell(ws, f"T{t}", fx("=SUM([H.mv])/Divisor"), F_MONEY2)
    _cell(ws, f"U{t}", f'=IF(N(T{t})=0,"",T{t}/T{t})', F_PCT)
    total_look(ws, t, 2, 21)
    grey_when_empty(ws, f"B{IS_FIRST}:B{t - 2}", IS_FIRST, ["T"])

    note(ws, f"B{t + 2}", "'% of theirs' is the category's share of that entity's own portfolio; "
         "on the Total row it is the entity's share of the family. Entities appear in the order "
         "of the Entities sheet.")
    put(ws, f"B{t + 3}", fx('=IF(COUNTA([E.id])>8,COUNTA([E.id])-8&" more entit(ies) are not '
                            'shown here but are in the Family column: see Checks.","")'),
        f=font(8, True, italic=True, color="C00000"), al=LEFT)

    w = {"A": 2, "B": 38, "S": SPACER, "T": 12, "U": 10}
    for v in pairs:
        w[v] = 11
        w[col(ci(v) + 1)] = 8
    widths(ws, w)
    ws.freeze_panes = f"C{IS_FIRST}"
    page_setup(ws, f"A1:U{t + 3}")
    return ws
