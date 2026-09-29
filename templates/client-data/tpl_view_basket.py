"""Basket Allocation (each basket's value by asset class, against the family's
own target) and Basket Detail (every holding in each basket, largest first).

The four baskets are Glow's own: Stable Growth, Entrepreneurial Growth,
Liquidity, and Thematic & Tactical. A security takes its category's default
basket from the Lists sheet unless the Securities sheet overrides it.
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

from tpl_schema import fx
from tpl_style import (BLUE, BOX, CENTER, F_DATE, F_MONEY2, F_PCT, INPUT_BG, LEFT, LEFT_WRAP,
                       NAVY, RIGHT, band, col, define, fill, font, put, widths)
from tpl_taxonomy import BASKETS, LIST_FIRST_ROW
from tpl_view_common import UNIT_LINE, note, page_setup, total_look, view_title
from tpl_view_pa import SPACER, _cell, _hdr

BA_FIRST = 7                       # first basket row on Basket Allocation
BA_TOTAL = BA_FIRST + len(BASKETS)  # 11
DEMO_TARGETS = [0.45, 0.12, 0.15, 0.28]

BD_SECTIONS = [5, 26, 47, 68]      # first row of each basket's section on Basket Detail
BD_SLOTS = 15
ABS_SHADE = "FCE4D6"


def _in_xirr_ratio(num_key, den_key, crit_formula):
    """A value-weighted return over the holdings that report one; blank where none do."""
    return (f'=IF(SUMIFS([H.{den_key}],{crit_formula})=0,"",'
            f'SUMIFS([H.{num_key}],{crit_formula})/SUMIFS([H.{den_key}],{crit_formula}))')


def build_basket_allocation(wb):
    ws = wb.create_sheet("Basket Allocation")
    view_title(ws, "Basket Allocation",
               "=" + UNIT_LINE + '&" · each basket\'s value by asset class, against the '
               'family\'s target"', "R1")

    for ref, text, to in [("B5", "Basket", "B6"), ("C5", "Code", "C6"),
                          ("D5", "Market value by asset class", "G5"),
                          ("H5", "Basket total", "H6"),
                          ("J5", "Actual allocation (%)", "J6"),
                          ("K5", "Target allocation (%)", "K6"),
                          ("L5", "Gap: actual − target (% points)", "L6"),
                          ("N5", "Investment at cost", "N6"),
                          ("O5", "Gain / (loss) on what is held", "O6"),
                          ("P5", "Absolute return (%)", "P6"),
                          ("Q5", "XIRR (%)", "Q6"),
                          ("R5", "Share of value with an XIRR", "R6")]:
        _hdr(ws, ref, text, to)
    for i, letter in enumerate("DEFG"):
        _hdr(ws, f"{letter}6", f"=Lists!$A${LIST_FIRST_ROW + i}")
    ws["K5"].fill = fill("7F6000")          # the one typed column: gold, not navy
    ws["K6"].fill = fill("7F6000")
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 30

    t = BA_TOTAL
    D = "Divisor"
    for i, (_code, _name, _desc) in enumerate(BASKETS):
        r = BA_FIRST + i
        k = f"[H.basket],$C{r}"
        put(ws, f"B{r}", f"=Lists!$L${LIST_FIRST_ROW + i}", f=font(9, True, color=NAVY), al=LEFT)
        put(ws, f"C{r}", f"=Lists!$K${LIST_FIRST_ROW + i}", f=font(9, True, color=NAVY),
            al=CENTER)
        for letter in "DEFG":
            _cell(ws, f"{letter}{r}", fx(f"=SUMIFS([H.mv],{k},[H.ac],{letter}$6)/{D}"), F_MONEY2)
        _cell(ws, f"H{r}", fx(f"=SUMIFS([H.mv],{k})/{D}"), F_MONEY2)
        ws[f"H{r}"].font = font(9, True)
        _cell(ws, f"J{r}", f'=IF(N($H${t})=0,"",H{r}/$H${t})', F_PCT)
        put(ws, f"K{r}", DEMO_TARGETS[i], f=font(9, True, color=BLUE), bg=INPUT_BG, fmt=F_PCT,
            al=RIGHT, border=BOX)
        _cell(ws, f"L{r}", f'=IF(OR(J{r}="",K{r}=""),"",J{r}-K{r})', F_PCT)
        _cell(ws, f"N{r}", fx(f"=SUMIFS([H.cost],{k})/{D}"), F_MONEY2)
        _cell(ws, f"O{r}", fx(f"=SUMIFS([H.ugain],{k})/{D}"), F_MONEY2)
        _cell(ws, f"P{r}", f'=IF(N(N{r})=0,"",O{r}/N{r})', F_PCT)
        _cell(ws, f"Q{r}", fx(_in_xirr_ratio("h_mx", "h_m", k)), F_PCT)
        _cell(ws, f"R{r}", fx(f'=IF(N(H{r})=0,"",SUMIFS([H.h_m],{k})/{D}/H{r})'), F_PCT)

    # Targets: a fraction between 0 and 1
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                        allow_blank=True)
    dv.error, dv.errorTitle = "Type the target as a percentage between 0% and 100%.", "Target"
    dv.prompt, dv.promptTitle, dv.showInputMessage = (
        "The family's target share of the portfolio for this basket. The four should add to "
        "100%.", "Target allocation", True)
    ws.add_data_validation(dv)
    dv.add(f"K{BA_FIRST}:K{t - 1}")

    # Total: straight from the register, so Checks can tie the rows to it
    put(ws, f"B{t}", "Total", f=font(9, True, color=NAVY), al=LEFT)
    for letter in "DEFG":
        _cell(ws, f"{letter}{t}", fx(f"=SUMIFS([H.mv],[H.ac],{letter}$6)/{D}"), F_MONEY2)
    _cell(ws, f"H{t}", fx(f"=SUM([H.mv])/{D}"), F_MONEY2)
    _cell(ws, f"J{t}", f"=SUM(J{BA_FIRST}:J{t - 1})", F_PCT)
    _cell(ws, f"K{t}", f"=SUM(K{BA_FIRST}:K{t - 1})", F_PCT)
    _cell(ws, f"N{t}", fx(f"=SUM([H.cost])/{D}"), F_MONEY2)
    _cell(ws, f"O{t}", fx(f"=SUM([H.ugain])/{D}"), F_MONEY2)
    _cell(ws, f"P{t}", f'=IF(N(N{t})=0,"",O{t}/N{t})', F_PCT)
    _cell(ws, f"Q{t}", fx('=IF(SUM([H.h_m])=0,"",SUM([H.h_mx])/SUM([H.h_m]))'), F_PCT)
    _cell(ws, f"R{t}", fx(f'=IF(N(H{t})=0,"",SUM([H.h_m])/{D}/H{t})'), F_PCT)
    total_look(ws, t, 2, 18)

    # The gap is red outside the family's rebalancing band, green inside it
    gap = f"L{BA_FIRST}:L{t - 1}"
    ws.conditional_formatting.add(gap, FormulaRule(
        formula=[f"AND(ISNUMBER($L{BA_FIRST}),ABS($L{BA_FIRST})>BasketTol)"],
        font=Font(name="Arial", size=9, bold=True, color="9C0006"),
        fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add(gap, FormulaRule(
        formula=[f"AND(ISNUMBER($L{BA_FIRST}),ABS($L{BA_FIRST})<=BasketTol)"],
        font=Font(name="Arial", size=9, bold=True, color="006100"),
        fill=PatternFill("solid", fgColor="C6EFCE")))
    # Targets that do not add to 100% are flagged on the total row
    ws.conditional_formatting.add(f"K{t}", FormulaRule(
        formula=[f"ABS($K${t}-1)>0.0001"],
        font=Font(name="Arial", size=9, bold=True, color="9C0006"),
        fill=PatternFill("solid", fgColor="FFC7CE")))

    tol = t + 2
    put(ws, f"B{tol}", "Rebalancing band (± % points)", f=font(9, True), al=LEFT, border=BOX)
    put(ws, f"C{tol}", 0.05, f=font(9, True, color=BLUE), bg=INPUT_BG, fmt=F_PCT, al=RIGHT,
        border=BOX)
    note(ws, f"D{tol}", "A gap inside the band is green; outside it, red. Type the band the family "
         "accepts before it rebalances.")
    define(wb, "BasketTol", f"'Basket Allocation'!$C${tol}")

    top = tol + 2
    band(ws, top, 2, 18, "WHAT EACH BASKET HOLDS — the family's own grouping (Lists sheet)",
         bg=NAVY, color="FFFFFF")
    for i in range(len(BASKETS)):
        r = top + 1 + i
        lr = LIST_FIRST_ROW + i
        put(ws, f"B{r}", f"=Lists!$L${lr}&\" (\"&Lists!$K${lr}&\")\"",
            f=font(9, True, color=NAVY), al=LEFT)
        put(ws, f"C{r}", f"=Lists!$M${lr}", f=font(9), al=LEFT)
    n0 = top + len(BASKETS) + 2
    note(ws, f"B{n0}", "A security takes its category's default basket (Lists sheet) unless its "
         "Securities row overrides it: a direct stock the family treats as core can sit in "
         "Stable Growth.")
    note(ws, f"B{n0 + 1}", "XIRR is value-weighted over the holdings held a year or more that "
         "report one; the last column says how much of the basket's value that covers.")
    note(ws, f"B{n0 + 2}", "Basket Detail lists every holding in each basket, largest first.")

    widths(ws, {"A": 2, "B": 30, "C": 9, "D": 12, "E": 12, "F": 12, "G": 12, "H": 13,
                "I": SPACER, "J": 11, "K": 11, "L": 12, "M": SPACER, "N": 13, "O": 13, "P": 11,
                "Q": 10, "R": 12})
    ws.freeze_panes = f"C{BA_FIRST}"
    page_setup(ws, f"A1:R{n0 + 2}")
    return ws


def build_basket_detail(wb):
    ws = wb.create_sheet("Basket Detail")
    view_title(ws, "Basket Detail",
               "=" + UNIT_LINE + '&" · every holding in each basket, largest first (up to '
               f'{BD_SLOTS} per basket; each total covers every holding)"', "O1")
    heads = [("B", "Security / product"), ("C", "Entity"), ("D", "Advisor / manager"),
             ("E", "Category"), ("F", "First investment"), ("G", "Investment at cost"),
             ("H", "Market value"), ("I", "% of basket"), ("J", "Gain / (loss) on what is held"),
             ("K", "Return shown"), ("L", "Return (%)"), ("M", "Benchmark return (%)"),
             ("N", "Alpha (% points)"), ("O", "Benchmark")]
    D = "Divisor"
    for i, s in enumerate(BD_SECTIONS):
        lr = LIST_FIRST_ROW + i
        # Section title, with the code the slots look up (hidden column R)
        put(ws, f"B{s}", f"=Lists!$L${lr}&\"  ·  \"&Lists!$K${lr}",
            f=font(10, True, color="FFFFFF"), bg=NAVY, al=LEFT)
        ws.merge_cells(f"B{s}:O{s}")
        for n in range(3, 16):
            ws[f"{col(n)}{s}"].fill = fill(NAVY)
        ws.row_dimensions[s].height = 20
        put(ws, f"R{s}", f"=Lists!$K${lr}", f=font(8, color="808080"))
        hdr = s + 1
        for letter, text in heads:
            _hdr(ws, f"{letter}{hdr}", text)
        ws.row_dimensions[hdr].height = 30
        first, tot = s + 2, s + 2 + BD_SLOTS
        code = f"$R${s}"
        for k in range(1, BD_SLOTS + 1):
            r = first + k - 1
            put(ws, f"S{r}", k, f=font(8, color="808080"))
            put(ws, f"Q{r}", fx(f'=IFERROR(MATCH({code}&"|"&$S{r},[H.h_key],0),"")'),
                f=font(8, color="808080"))
            q = f"$Q{r}"

            def txt(key):
                return fx(f'=IF({q}="","",INDEX([H.{key}],{q})&"")')

            def num(key, divide=False):
                v = f"INDEX([H.{key}],{q})"
                out = f"{v}/{D}" if divide else v
                return fx(f'=IF({q}="","",IF({v}="","",{out}))')

            _cell(ws, f"B{r}", txt("name"), None, LEFT)
            _cell(ws, f"C{r}", txt("ent"), None, LEFT)
            _cell(ws, f"D{r}", txt("adv"), None, LEFT)
            _cell(ws, f"E{r}", txt("cat"), None, LEFT)
            _cell(ws, f"F{r}", num("first"), F_DATE, CENTER)
            _cell(ws, f"G{r}", num("cost", True), F_MONEY2)
            _cell(ws, f"H{r}", num("mv", True), F_MONEY2)
            _cell(ws, f"I{r}", f'=IF(OR({q}="",N($H${tot})=0),"",H{r}/$H${tot})', F_PCT)
            _cell(ws, f"J{r}", num("ugain", True), F_MONEY2)
            _cell(ws, f"K{r}", txt("shown"), None, CENTER)
            ret = (f'=IF({q}="","",IF(K{r}="XIRR",INDEX([H.xirr],{q}),IF(K{r}="Absolute",'
                   f'IF(INDEX([H.absret],{q})="","",INDEX([H.absret],{q})),"")))')
            _cell(ws, f"L{r}", fx(ret), F_PCT)
            bret = (f'=IF({q}="","",IF(K{r}="XIRR",IF(INDEX([H.bxirr],{q})="","",'
                    f'INDEX([H.bxirr],{q})),IF(K{r}="Absolute",IF(INDEX([H.babs],{q})="","",'
                    f'INDEX([H.babs],{q})),"")))')
            _cell(ws, f"M{r}", fx(bret), F_PCT)
            _cell(ws, f"N{r}", f'=IF(OR(L{r}="",M{r}=""),"",L{r}-M{r})', F_PCT)
            _cell(ws, f"O{r}", txt("bench"), None, LEFT)
        # Shade the rows that show an absolute return (held under a year)
        ws.conditional_formatting.add(f"B{first}:O{tot - 1}", FormulaRule(
            formula=[f'$K{first}="Absolute"'], fill=PatternFill("solid", fgColor=ABS_SHADE)))

        # Total: every holding in the basket, not only the ones listed
        crit = f"[H.basket],{code}"
        put(ws, f"B{tot}", f"=\"Total — \"&Lists!$L${lr}", f=font(9, True, color=NAVY), al=LEFT)
        _cell(ws, f"G{tot}", fx(f"=SUMIFS([H.cost],{crit})/{D}"), F_MONEY2)
        _cell(ws, f"H{tot}", fx(f"=SUMIFS([H.mv],{crit})/{D}"), F_MONEY2)
        _cell(ws, f"I{tot}", f'=IF(N(H{tot})=0,"",SUM(I{first}:I{tot - 1}))', F_PCT)
        _cell(ws, f"J{tot}", fx(f"=SUMIFS([H.ugain],{crit})/{D}"), F_MONEY2)
        # The basket's XIRR where any holding reports one; otherwise its absolute
        # return on cost, against the cost-weighted benchmark (Portfolio Allocation's rule)
        _cell(ws, f"K{tot}", fx(f'=IF(SUMIFS([H.h_m],{crit})>0,"XIRR",'
                                f'IF(N(G{tot})>0,"Absolute",""))'), None, CENTER)
        xr = _in_xirr_ratio("h_mx", "h_m", crit)[1:]
        bx = _in_xirr_ratio("h_mbx", "h_m", crit)[1:]
        ba = _in_xirr_ratio("h_mba", "h_mbam", crit)[1:]
        _cell(ws, f"L{tot}", fx(f'=IF(K{tot}="XIRR",{xr},IF(K{tot}="Absolute",J{tot}/G{tot},""))'),
              F_PCT)
        _cell(ws, f"M{tot}", fx(f'=IF(K{tot}="XIRR",{bx},IF(K{tot}="Absolute",{ba},""))'), F_PCT)
        _cell(ws, f"N{tot}", f'=IF(OR(L{tot}="",M{tot}=""),"",L{tot}-M{tot})', F_PCT)
        total_look(ws, tot, 2, 15)
        count = f"SUMPRODUCT(([H.basket]={code})*ISNUMBER([H.mv]))"
        put(ws, f"B{tot + 1}", fx(f'=IF({count}>{BD_SLOTS},{count}-{BD_SLOTS}&" more holding(s) '
                                  f'in this basket are not listed; the total includes them.",'
                                  f'IF({count}=0,"Nothing is held in this basket.",""))'),
            f=font(8, True, italic=True, color="C00000"), al=LEFT)
        if i:
            ws.row_breaks.append(Break(id=s - 1))

    leg = BD_SECTIONS[-1] + BD_SLOTS + 5
    put(ws, f"B{leg}", "Held under a year", f=font(8, True), bg=ABS_SHADE, al=CENTER, border=BOX)
    note(ws, f"C{leg}", "Shaded rows show the absolute return since the first investment; the "
         "others the XIRR the statement reports, against the benchmark's return over the same "
         "period.")
    note(ws, f"C{leg + 1}", "A total row shows the basket's XIRR, value-weighted over the holdings "
         "that report one; where none does, its absolute return on cost. 'At cost' holdings "
         "(private equity, pre-IPO) carry no return until a priced round or a listing, and a "
         "holding with no investment date (a bank balance) shows none.")
    for letter in "QRS":
        ws.column_dimensions[letter].hidden = True
    widths(ws, {"A": 2, "B": 38, "C": 17, "D": 22, "E": 26, "F": 11, "G": 12, "H": 12, "I": 9,
                "J": 12, "K": 10, "L": 9, "M": 10, "N": 9, "O": 26, "P": 2})
    ws.freeze_panes = "C5"
    page_setup(ws, f"A1:O{leg + 1}")
    return ws
